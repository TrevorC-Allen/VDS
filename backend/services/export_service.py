"""Generate shareable workbench export artifacts from stable response fields."""

from __future__ import annotations

import base64
import csv
from datetime import datetime, timezone
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from typing import Any
from xml.sax.saxutils import escape as xml_escape
import zipfile

from backend.schemas.data_agent_schema import to_json_ready


BLOCKED_EXPORT_KEYS = {
    "activity_trace_v2",
    "debug",
    "execution_artifacts",
    "reasoning_trace_view",
    "raw_prompt",
    "trace",
    "scorer",
    "".join(("standard", "_answer")),
}


def generate_export_artifacts(response: dict[str, Any], *, runs_root: str | Path) -> dict[str, Any]:
    """Write export files and return a manifest safe for frontend download menus."""

    run_id = _safe_run_id(str(response.get("run_id") or ""))
    if not run_id:
        return _manifest("", unavailable=[_unavailable("summary", "all", "run_id is missing.")])
    exports_dir = Path(runs_root) / run_id / "exports"
    exports_dir.mkdir(parents=True, exist_ok=True)
    artifacts: list[dict[str, Any]] = []
    unavailable: list[dict[str, Any]] = []
    stable = _stable_summary_payload(response)

    result = response.get("result") if isinstance(response.get("result"), dict) else {}
    columns = [str(column) for column in result.get("columns") or []]
    rows = [row for row in result.get("rows") or [] if isinstance(row, dict)]
    if rows and columns:
        artifacts.append(_write_result_csv(exports_dir, run_id, columns, rows))
        try:
            artifacts.append(_write_result_xlsx(exports_dir, run_id, columns, rows))
        except Exception as exc:  # noqa: BLE001 - a missing optional writer should not break the response.
            unavailable.append(_unavailable("result_table", "xlsx", str(exc)))
    else:
        unavailable.append(_unavailable("result_table", "csv", "结果为空或不是行列表。"))
        unavailable.append(_unavailable("result_table", "xlsx", "结果为空或不是行列表。"))

    chart = response.get("chart") if isinstance(response.get("chart"), dict) else {}
    chart_artifact = _write_chart_fallback(exports_dir, run_id, chart)
    if chart_artifact is not None:
        artifacts.append(chart_artifact)
    elif chart:
        unavailable.append(_unavailable("chart", "server_file", "当前图表由前端交互 SVG 渲染，下载时由浏览器序列化。"))
    else:
        unavailable.append(_unavailable("chart", "server_file", "本次结果没有可下载图表。"))

    try:
        artifacts.append(_write_summary_xlsx(exports_dir, run_id, stable, columns, rows))
    except Exception as exc:  # noqa: BLE001
        unavailable.append(_unavailable("summary_report", "xlsx", str(exc)))
    try:
        artifacts.append(_write_summary_pptx(exports_dir, run_id, stable))
    except Exception as exc:  # noqa: BLE001
        unavailable.append(_unavailable("summary_report", "pptx", str(exc)))
    try:
        artifacts.append(_write_summary_pdf(exports_dir, run_id, stable))
    except Exception as exc:  # noqa: BLE001
        unavailable.append(_unavailable("summary_report", "pdf", str(exc)))

    manifest = _manifest(run_id, artifacts=artifacts, unavailable=unavailable)
    manifest_path = exports_dir / "manifest.json"
    manifest_path.write_text(json.dumps(to_json_ready(manifest), ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


def resolve_export_artifact(runs_root: str | Path, run_id: str, artifact_id: str) -> tuple[Path, dict[str, Any]] | None:
    safe_run_id = _safe_run_id(run_id)
    safe_artifact_id = _safe_artifact_id(artifact_id)
    if not safe_run_id or not safe_artifact_id:
        return None
    manifest_path = Path(runs_root) / safe_run_id / "exports" / "manifest.json"
    if not manifest_path.exists():
        return None
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    for artifact in manifest.get("artifacts") or []:
        if not isinstance(artifact, dict) or artifact.get("artifact_id") != safe_artifact_id:
            continue
        file_name = str(artifact.get("file_name") or "")
        if not file_name or "/" in file_name or "\\" in file_name:
            return None
        path = manifest_path.parent / file_name
        if path.exists():
            return path, artifact
    return None


def _stable_summary_payload(response: dict[str, Any]) -> dict[str, Any]:
    result = response.get("result") if isinstance(response.get("result"), dict) else {}
    chart = response.get("chart") if isinstance(response.get("chart"), dict) else {}
    insight = response.get("insight") if isinstance(response.get("insight"), dict) else {}
    process = response.get("process_view_v2") if isinstance(response.get("process_view_v2"), dict) else {}
    sources = response.get("source_references") if isinstance(response.get("source_references"), list) else []
    warnings = [str(item) for item in response.get("warnings") or [] if str(item).strip()]
    caveats = [str(item) for item in insight.get("caveats") or [] if str(item).strip()]
    payload = {
        "question": str(response.get("question") or ""),
        "answer": str(response.get("answer") or ""),
        "answer_type": str(response.get("answer_type") or ""),
        "dataset_id": str(response.get("dataset_id") or ""),
        "run_id": str(response.get("run_id") or ""),
        "result": {
            "columns": list(result.get("columns") or []),
            "rows": list(result.get("rows") or [])[:200],
            "value": result.get("value"),
        },
        "chart": {
            "chart_type": chart.get("chart_type"),
            "title": chart.get("title"),
            "x": chart.get("x"),
            "y": chart.get("y"),
            "reason": chart.get("reason"),
        }
        if chart
        else None,
        "insight": {
            "summary": insight.get("summary"),
            "next_step": insight.get("next_step"),
            "caveats": caveats,
        }
        if insight
        else None,
        "sources": sources[:20],
        "process_summary": process.get("summary") or "",
        "process_steps": [
            {
                "title": step.get("title"),
                "summary": step.get("summary"),
                "status": step.get("status"),
            }
            for step in (process.get("steps") or [])[:12]
            if isinstance(step, dict)
        ],
        "caveats": caveats + warnings,
        "created_at": _now_iso(),
    }
    return _strip_blocked_payload(to_json_ready(payload))


def _strip_blocked_payload(value: Any) -> Any:
    if isinstance(value, dict):
        cleaned: dict[str, Any] = {}
        for key, item in value.items():
            lowered = str(key).lower()
            if any(blocked in lowered for blocked in BLOCKED_EXPORT_KEYS):
                continue
            cleaned[key] = _strip_blocked_payload(item)
        return cleaned
    if isinstance(value, list):
        return [_strip_blocked_payload(item) for item in value]
    if isinstance(value, str):
        lowered = value.lower()
        if any(blocked in lowered for blocked in ("chain_of_thought", "raw prompt", "standard answer", "scorer")):
            return ""
    return value


def _write_result_csv(exports_dir: Path, run_id: str, columns: list[str], rows: list[dict[str, Any]]) -> dict[str, Any]:
    path = exports_dir / "result_table.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow([_safe_spreadsheet_cell(column) for column in columns])
        for row in rows:
            writer.writerow([_safe_spreadsheet_cell(row.get(column)) for column in columns])
    return _artifact(run_id, "result_table_csv", "result_table", "csv", "结果表 CSV", "text/csv; charset=utf-8", path.name)


def _write_result_xlsx(exports_dir: Path, run_id: str, columns: list[str], rows: list[dict[str, Any]]) -> dict[str, Any]:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    path = exports_dir / "result_table.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Result"
    sheet.append([_safe_spreadsheet_cell(column) for column in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([_safe_spreadsheet_cell(row.get(column)) for column in columns])
    for column_cells in sheet.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 10), 42)
    workbook.save(path)
    return _artifact(run_id, "result_table_xlsx", "result_table", "xlsx", "结果表 Excel", _xlsx_mime(), path.name)


def _write_summary_xlsx(
    exports_dir: Path,
    run_id: str,
    stable: dict[str, Any],
    columns: list[str],
    rows: list[dict[str, Any]],
) -> dict[str, Any]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    path = exports_dir / "summary.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Field", "Value"])
    summary.append(["Question", _safe_spreadsheet_cell(stable.get("question") or "")])
    summary.append(["Answer", _safe_spreadsheet_cell(stable.get("answer") or "")])
    summary.append(["Insight", _safe_spreadsheet_cell((stable.get("insight") or {}).get("summary") or "")])
    summary.append(["Process", _safe_spreadsheet_cell(stable.get("process_summary") or "")])
    summary.append(["Caveats", _safe_spreadsheet_cell("\n".join(stable.get("caveats") or []))])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    for row in summary.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    summary.column_dimensions["A"].width = 18
    summary.column_dimensions["B"].width = 88

    result_sheet = workbook.create_sheet("Result")
    if columns and rows:
        result_sheet.append([_safe_spreadsheet_cell(column) for column in columns])
        for cell in result_sheet[1]:
            cell.font = Font(bold=True)
        for row in rows[:1000]:
            result_sheet.append([_safe_spreadsheet_cell(row.get(column)) for column in columns])
    else:
        result_sheet.append(["Unavailable"])
        result_sheet.append(["本次没有可导出的行列表结果。"])

    sources_sheet = workbook.create_sheet("Sources")
    sources_sheet.append(["File", "Tables", "Rows"])
    for source in stable.get("sources") or []:
        if not isinstance(source, dict):
            continue
        sources_sheet.append(
            [
                _safe_spreadsheet_cell(source.get("file_name") or ""),
                _safe_spreadsheet_cell(
                    ", ".join(str(table.get("table_name") or "") for table in source.get("tables") or [] if isinstance(table, dict))
                ),
                source.get("row_count") or "",
            ]
        )
    workbook.save(path)
    return _artifact(run_id, "summary_xlsx", "summary_report", "xlsx", "Excel 摘要", _xlsx_mime(), path.name)


def _write_summary_pptx(exports_dir: Path, run_id: str, stable: dict[str, Any]) -> dict[str, Any]:
    path = exports_dir / "summary.pptx"
    title = _clip(stable.get("question") or "VDS Summary", 80)
    lines = [
        _clip(stable.get("answer") or "", 220),
        _clip((stable.get("insight") or {}).get("summary") or "", 160),
        _clip(stable.get("process_summary") or "", 160),
    ]
    lines = [line for line in lines if line]
    _write_minimal_pptx(path, title=title, lines=lines or ["No summary text available."])
    return _artifact(
        run_id,
        "summary_pptx",
        "summary_report",
        "pptx",
        "PPT 摘要",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        path.name,
    )


def _write_summary_pdf(exports_dir: Path, run_id: str, stable: dict[str, Any]) -> dict[str, Any]:
    path = exports_dir / "summary.pdf"
    lines = [
        "VDS Summary Report",
        "",
        f"Question: {stable.get('question') or ''}",
        f"Answer: {stable.get('answer') or ''}",
        f"Insight: {(stable.get('insight') or {}).get('summary') or ''}",
        f"Process: {stable.get('process_summary') or ''}",
    ]
    if not _write_pdf_with_soffice(path, lines):
        if not _write_pdf_with_reportlab(path, lines):
            _write_minimal_pdf(path, lines)
    return _artifact(run_id, "summary_pdf", "summary_report", "pdf", "PDF 摘要", "application/pdf", path.name)


def _write_chart_fallback(exports_dir: Path, run_id: str, chart: dict[str, Any]) -> dict[str, Any] | None:
    data_uri = str(chart.get("image_data_uri") or "")
    if data_uri.startswith("data:image/"):
        header, _, encoded = data_uri.partition(",")
        match = re.match(r"data:image/([^;,]+)(?:;[^,]*)*;base64", header, re.IGNORECASE)
        if match and encoded:
            extension = _image_extension_from_mime_subtype(match.group(1))
            if extension not in {"png", "jpg", "jpeg", "svg", "webp"}:
                extension = "png"
            path = exports_dir / f"chart.{extension}"
            path.write_bytes(base64.b64decode(encoded))
            if extension == "svg":
                mime_type = "image/svg+xml"
            else:
                mime_type = f"image/{'jpeg' if extension in {'jpg', 'jpeg'} else extension}"
            return _artifact(run_id, f"chart_{extension}", "chart", extension, f"图表 {extension.upper()}", mime_type, path.name)
    svg = str(chart.get("svg") or "")
    if svg.strip().startswith("<svg"):
        path = exports_dir / "chart.svg"
        path.write_text(svg, encoding="utf-8")
        return _artifact(run_id, "chart_svg", "chart", "svg", "图表 SVG", "image/svg+xml", path.name)
    return None


def _write_pdf_with_soffice(path: Path, lines: list[str]) -> bool:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return False
    html = _summary_html(lines)
    with tempfile.TemporaryDirectory() as temp_dir:
        temp_path = Path(temp_dir)
        html_path = temp_path / "summary.html"
        html_path.write_text(html, encoding="utf-8")
        try:
            completed = subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(temp_path), str(html_path)],
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                timeout=20,
                check=False,
            )
        except Exception:
            return False
        pdf_path = temp_path / "summary.pdf"
        if completed.returncode == 0 and pdf_path.exists() and pdf_path.stat().st_size > 0:
            shutil.copy2(pdf_path, path)
            return True
    return False


def _summary_html(lines: list[str]) -> str:
    body = "".join(f"<p>{xml_escape(str(line))}</p>" if line else "<br/>" for line in lines)
    return (
        "<!doctype html><html><head><meta charset='utf-8'>"
        "<style>body{font-family:-apple-system,BlinkMacSystemFont,'PingFang SC',Arial,sans-serif;font-size:14px;line-height:1.55;}"
        "p{margin:0 0 10px;white-space:pre-wrap;}</style></head><body>"
        f"{body}</body></html>"
    )


def _write_pdf_with_reportlab(path: Path, lines: list[str]) -> bool:
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfgen import canvas
    except Exception:
        return False
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        _, page_height = letter
        doc = canvas.Canvas(str(path), pagesize=letter)
        doc.setTitle("VDS Summary Report")
        y = page_height - 52
        for line in lines[:60]:
            wrapped = _wrap_pdf_line(str(line), limit=88) if line else [""]
            for segment in wrapped:
                if y < 52:
                    doc.showPage()
                    y = page_height - 52
                doc.setFont("STSong-Light", 12)
                doc.drawString(52, y, segment)
                y -= 18
            if line:
                y -= 4
        doc.save()
        return path.exists() and path.stat().st_size > 0
    except Exception:
        return False


def _write_minimal_pdf(path: Path, lines: list[str]) -> None:
    safe_lines = [_pdf_ascii(line) for line in lines[:40]]
    text_ops = ["BT", "/F1 12 Tf", "50 760 Td"]
    for index, line in enumerate(safe_lines):
        if index:
            text_ops.append("0 -18 Td")
        text_ops.append(f"({_pdf_escape(line[:110])}) Tj")
    text_ops.append("ET")
    stream = "\n".join(text_ops).encode("latin-1", errors="replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    content = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(content))
        content.extend(f"{index} 0 obj\n".encode("ascii"))
        content.extend(obj)
        content.extend(b"\nendobj\n")
    xref_offset = len(content)
    content.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        content.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    content.extend(f"trailer << /Root 1 0 R /Size {len(objects) + 1} >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii"))
    path.write_bytes(bytes(content))


def _write_minimal_pptx(path: Path, *, title: str, lines: list[str]) -> None:
    paragraph_xml = "".join(
        f"<a:p><a:r><a:rPr lang=\"zh-CN\" sz=\"2000\"/><a:t>{xml_escape(line)}</a:t></a:r></a:p>"
        for line in lines[:8]
    )
    slide_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <p:cSld><p:spTree>
    <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>
    <p:sp><p:nvSpPr><p:cNvPr id="2" name="Title"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="685800" y="457200"/><a:ext cx="7772400" cy="914400"/></a:xfrm></p:spPr><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="zh-CN" sz="3200" b="1"/><a:t>{xml_escape(title)}</a:t></a:r></a:p></p:txBody></p:sp>
    <p:sp><p:nvSpPr><p:cNvPr id="3" name="Summary"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr><a:xfrm><a:off x="685800" y="1500000"/><a:ext cx="7772400" cy="4300000"/></a:xfrm></p:spPr><p:txBody><a:bodyPr wrap="square"/><a:lstStyle/>{paragraph_xml}</p:txBody></p:sp>
  </p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr>
</p:sld>"""
    presentation_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:presentation xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <p:sldIdLst><p:sldId id="256" r:id="rId1"/></p:sldIdLst>
  <p:sldSz cx="9144000" cy="6858000" type="screen4x3"/><p:notesSz cx="6858000" cy="9144000"/>
</p:presentation>"""
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>
  <Override PartName="/ppt/slides/slide1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>
</Types>""",
        )
        archive.writestr(
            "_rels/.rels",
            """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/>
</Relationships>""",
        )
        archive.writestr(
            "ppt/_rels/presentation.xml.rels",
            """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" Target="slides/slide1.xml"/>
</Relationships>""",
        )
        archive.writestr("ppt/presentation.xml", presentation_xml)
        archive.writestr("ppt/slides/slide1.xml", slide_xml)


def _artifact(
    run_id: str,
    artifact_id: str,
    artifact_type: str,
    fmt: str,
    display_name: str,
    mime_type: str,
    file_name: str,
) -> dict[str, Any]:
    return {
        "artifact_id": artifact_id,
        "artifact_type": artifact_type,
        "format": fmt,
        "display_name": display_name,
        "mime_type": mime_type,
        "file_name": file_name,
        "download_name": file_name,
        "created_at": _now_iso(),
        "download_url": f"/api/data-agent/runs/{run_id}/exports/{artifact_id}",
        "source_run_id": run_id,
    }


def _manifest(
    run_id: str,
    *,
    artifacts: list[dict[str, Any]] | None = None,
    unavailable: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    return {
        "version": "export_artifacts.v1",
        "source_run_id": run_id,
        "artifacts": artifacts or [],
        "unavailable": unavailable or [],
        "created_at": _now_iso(),
    }


def _unavailable(artifact_type: str, fmt: str, reason: str) -> dict[str, str]:
    return {"artifact_type": artifact_type, "format": fmt, "reason": reason}


def _safe_spreadsheet_cell(value: Any) -> Any:
    if isinstance(value, str) and _looks_like_spreadsheet_formula(value):
        return "'" + value
    return value


def _looks_like_spreadsheet_formula(value: str) -> bool:
    text = str(value or "")
    return bool(text) and (text[0] in {"=", "+", "-", "@", "\t", "\r", "\n"} or text.lstrip()[:1] in {"=", "+", "-", "@"})


def _xlsx_mime() -> str:
    return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _image_extension_from_mime_subtype(value: str) -> str:
    subtype = str(value or "").strip().lower()
    if subtype in {"jpeg", "jpg", "pjpeg"}:
        return "jpg"
    if subtype in {"svg", "svg+xml"}:
        return "svg"
    return subtype


def _safe_run_id(value: str) -> str:
    text = str(value or "").strip()
    return text if re.fullmatch(r"[A-Za-z0-9_-]{1,96}", text) else ""


def _safe_artifact_id(value: str) -> str:
    text = str(value or "").strip()
    return text if re.fullmatch(r"[A-Za-z0-9_-]{1,96}", text) else ""


def _clip(value: Any, limit: int) -> str:
    text = str(value or "").strip()
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _wrap_pdf_line(value: str, *, limit: int) -> list[str]:
    text = str(value or "").replace("\n", " ")
    return [text[index : index + limit] for index in range(0, len(text), limit)] or [""]


def _pdf_ascii(value: Any) -> str:
    text = str(value or "").replace("\n", " ")
    return text.encode("latin-1", errors="replace").decode("latin-1")


def _pdf_escape(value: str) -> str:
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()
