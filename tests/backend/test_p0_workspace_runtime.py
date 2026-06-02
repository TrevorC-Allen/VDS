"""P0 workspace persistence, run status, and export artifact contracts."""

from __future__ import annotations

import csv
import base64
import json
import tempfile
import unittest
import zipfile
from pathlib import Path

import backend.routers.data_agent as data_agent_router
from backend.services.data_agent_service import DataAgentService
from backend.services.export_service import generate_export_artifacts
from backend.storage.project_store import ProjectStore
from backend.storage.run_store import RunStore
from backend.storage.temp_file_store import TempFileStore
from data_agent_core.llm.client import MockLLMClient


class P0WorkspaceRuntimeTest(unittest.TestCase):
    def test_project_instructions_and_sources_are_versioned(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            store = ProjectStore(Path(temp_dir) / "projects")
            project = store.create_project(name="BI", instructions="利润 = 收入 - 成本")
            updated = store.update_project(project["project_id"], instructions="利润 = 收入 - 成本 - 退款")
            source = store.add_source(
                project["project_id"],
                source_type="saved_response",
                title="保存的回答",
                content="这是最终回答。",
                metadata={"conversation_id": "conv_demo", "run_id": "run_demo"},
            )

        self.assertEqual(2, updated["instructions_version"])
        self.assertTrue(updated["instructions_updated_at"])
        self.assertEqual(64, len(updated["content_hash"]))
        self.assertEqual("saved_response", source["source_type"])
        self.assertEqual(1, source["source_version"])
        self.assertEqual("active", source["source_status"])
        self.assertEqual(64, len(source["content_hash"]))

    def test_dataset_profile_restores_persisted_upload_after_restart(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            csv_path = root / "sales.csv"
            csv_path.write_text("city,sales\n上海,100\n北京,120\n", encoding="utf-8")
            first_service = DataAgentService(file_store=TempFileStore(root / "storage"), llm_client=MockLLMClient())
            upload = first_service.upload_dataset(csv_path)
            dataset_id = upload["dataset_id"]

            restarted_service = DataAgentService(file_store=TempFileStore(root / "storage"), llm_client=MockLLMClient())
            profile = restarted_service.get_dataset_profile(dataset_id)

        self.assertTrue(profile["success"])
        self.assertEqual(dataset_id, profile["dataset_id"])
        self.assertTrue(profile["can_analyze"])
        self.assertTrue(profile["restored_from_disk"])
        self.assertEqual("", profile["restore_error"])
        self.assertTrue(profile["source_files"])

    def test_dataset_profile_reports_reupload_when_sources_are_missing(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            csv_path = root / "sales.csv"
            csv_path.write_text("city,sales\n上海,100\n", encoding="utf-8")
            first_service = DataAgentService(file_store=TempFileStore(root / "storage"), llm_client=MockLLMClient())
            upload = first_service.upload_dataset(csv_path)
            dataset_id = upload["dataset_id"]
            for path in (root / "storage" / "datasets" / dataset_id / "source_file").glob("*"):
                path.unlink()

            restarted_service = DataAgentService(file_store=TempFileStore(root / "storage"), llm_client=MockLLMClient())
            profile = restarted_service.get_dataset_profile(dataset_id)

        self.assertTrue(profile["success"])
        self.assertFalse(profile["can_analyze"])
        self.assertEqual("needs_reupload", profile["source_status"])
        self.assertTrue(profile["restore_error"])

    def test_run_store_status_cancel_and_result_contract(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            store = RunStore(Path(temp_dir) / "runs")
            created = store.create_run(
                "run_demo",
                request={"question": "哪个城市最高？"},
                project_id="proj_demo",
                conversation_id="conv_demo",
                dataset_id="ds_demo",
            )
            cancelled = store.request_cancel("run_demo")
            cancel_requested = store.is_cancel_requested("run_demo")
            store.write_result("run_demo", {"success": True, "answer": "上海"})
            result = store.get_result("run_demo")

            self.assertEqual("queued", created["status"])
            self.assertEqual("cancel_requested", cancelled["status"])
            self.assertTrue(cancel_requested)
            self.assertEqual("上海", result["answer"])

    def test_run_store_retry_tracks_attempt_and_failure_category(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            store = RunStore(Path(temp_dir) / "runs")
            store.create_run("run_original", request={"question": "恢复数据集"}, dataset_id="ds_missing")
            retry = store.create_run("run_retry", request={"question": "恢复数据集"}, dataset_id="ds_missing", retry_of="run_original")

        self.assertEqual("run_original", retry["retry_of"])
        self.assertEqual(2, retry["attempt"])
        self.assertEqual(
            "dataset_restore",
            data_agent_router._failure_category_from_error(
                "Dataset not found in temporary store: ds_missing",
                failed_step="analyze_dataset",
                error_type="FILE_PARSE_ERROR",
            ),
        )

    def test_router_run_store_follows_active_service_storage_root(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            original_service = data_agent_router.service
            original_run_store = data_agent_router.run_store
            try:
                root = Path(temp_dir)
                data_agent_router.service = DataAgentService(
                    file_store=TempFileStore(root / "active_storage"),
                    llm_client=MockLLMClient(),
                )
                data_agent_router.run_store = RunStore(root / "stale_runs")

                store = data_agent_router._current_run_store()

                self.assertEqual(root / "active_storage" / "runs", store.root)
            finally:
                data_agent_router.service = original_service
                data_agent_router.run_store = original_run_store

    def test_message_response_exposes_assistant_message_id_for_saved_response_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            service = DataAgentService(
                file_store=TempFileStore(Path(temp_dir) / "storage"),
                llm_client=MockLLMClient(),
            )
            response = service.respond_to_message(question="你好")
            conversation = service.get_conversation(response["conversation_id"])["conversation"]
            assistant = next(message for message in conversation["messages"] if message["role"] == "assistant")

        self.assertTrue(response["message_id"].startswith("msg_"))
        self.assertEqual(assistant["message_id"], response["message_id"])

    def test_export_artifacts_include_tables_and_reports_without_debug_trace(self) -> None:
        response = {
            "success": True,
            "run_id": "run_export_demo",
            "dataset_id": "ds_demo",
            "question": "哪个城市销售额最高？",
            "answer": "上海销售额最高，为 120。",
            "result": {"columns": ["city", "sales"], "rows": [{"city": "上海", "sales": 120}]},
            "chart": {"chart_type": "bar", "title": "城市销售额", "x": "city", "y": "sales"},
            "insight": {"summary": "上海领先。", "caveats": ["样本量较小。"]},
            "source_references": [{"file_name": "sales.csv", "tables": [{"table_name": "sales"}], "row_count": 1}],
            "process_view_v2": {"summary": "读取数据、计算排序、整理回答。", "steps": [{"title": "执行分析", "summary": "完成排序。"}]},
            "debug": {"raw_prompt": "secret"},
            "activity_trace_v2": [{"summary": "internal"}],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            manifest = generate_export_artifacts(response, runs_root=Path(temp_dir) / "runs")
            exports_dir = Path(temp_dir) / "runs" / "run_export_demo" / "exports"
            csv_path = exports_dir / "result_table.csv"
            summary_path = exports_dir / "summary.xlsx"
            ppt_path = exports_dir / "summary.pptx"
            pdf_path = exports_dir / "summary.pdf"
            manifest_text = (exports_dir / "manifest.json").read_text(encoding="utf-8")
            with csv_path.open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.DictReader(handle))
            with zipfile.ZipFile(ppt_path) as archive:
                ppt_text = "\n".join(
                    archive.read(name).decode("utf-8", errors="ignore")
                    for name in archive.namelist()
                    if name.endswith(".xml")
                )

            formats = {artifact["format"] for artifact in manifest["artifacts"]}
            self.assertIn("csv", formats)
            self.assertIn("xlsx", formats)
            self.assertIn("pptx", formats)
            self.assertIn("pdf", formats)
            self.assertTrue(all(artifact.get("download_name") for artifact in manifest["artifacts"]))
            self.assertEqual("上海", rows[0]["city"])
            self.assertTrue(summary_path.exists())
            self.assertGreater(pdf_path.stat().st_size, 0)
            self.assertNotIn("raw_prompt", manifest_text)
            self.assertNotIn("activity_trace_v2", manifest_text)
            self.assertNotIn("secret", manifest_text)
            self.assertNotIn("raw_prompt", ppt_text)
            self.assertNotIn("secret", ppt_text)

    def test_export_spreadsheets_escape_formula_like_user_content(self) -> None:
        from openpyxl import load_workbook

        response = {
            "success": True,
            "run_id": "run_export_formula_escape",
            "dataset_id": "ds_demo",
            "question": '=HYPERLINK("http://example.test","q")',
            "answer": "+answer should be text",
            "result": {
                "columns": ["=city", "+sales", "@note"],
                "rows": [{"=city": "上海", "+sales": 120, "@note": " \t=cmd"}],
            },
            "insight": {"summary": "-insight should be text", "caveats": ["=caveat"]},
            "source_references": [{"file_name": "=sales.csv", "tables": [{"table_name": "+sales"}], "row_count": 1}],
            "process_view_v2": {"summary": "@process should be text"},
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            generate_export_artifacts(response, runs_root=Path(temp_dir) / "runs")
            exports_dir = Path(temp_dir) / "runs" / "run_export_formula_escape" / "exports"
            with (exports_dir / "result_table.csv").open(encoding="utf-8-sig", newline="") as handle:
                csv_rows = list(csv.reader(handle))
            result_workbook = load_workbook(exports_dir / "result_table.xlsx", data_only=False)
            summary_workbook = load_workbook(exports_dir / "summary.xlsx", data_only=False)

        self.assertEqual(["'=city", "'+sales", "'@note"], csv_rows[0])
        self.assertEqual("' \t=cmd", csv_rows[1][2])
        self.assertEqual(["'=city", "'+sales", "'@note"], [result_workbook["Result"][f"{column}1"].value for column in "ABC"])
        self.assertEqual("' \t=cmd", result_workbook["Result"]["C2"].value)
        self.assertEqual(
            ['\'=HYPERLINK("http://example.test","q")', "'+answer should be text", "'-insight should be text", "'@process should be text", "'=caveat"],
            [summary_workbook["Summary"][f"B{row}"].value for row in range(2, 7)],
        )
        self.assertEqual(["'=sales.csv", "'+sales", 1], [summary_workbook["Sources"][f"{column}2"].value for column in "ABC"])
        self.assertEqual(["'=city", "'+sales", "'@note"], [summary_workbook["Result"][f"{column}1"].value for column in "ABC"])

    def test_export_chart_svg_data_uri_keeps_svg_format(self) -> None:
        svg = "<svg xmlns=\"http://www.w3.org/2000/svg\" viewBox=\"0 0 120 80\"><rect width=\"120\" height=\"80\" fill=\"#fff\"/></svg>"
        response = {
            "success": True,
            "run_id": "run_export_chart_svg",
            "question": "生成图表",
            "answer": "已生成图表。",
            "chart": {"image_data_uri": "data:image/svg+xml;base64," + base64.b64encode(svg.encode("utf-8")).decode("ascii")},
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            manifest = generate_export_artifacts(response, runs_root=Path(temp_dir) / "runs")
            exports_dir = Path(temp_dir) / "runs" / "run_export_chart_svg" / "exports"
            chart_path = exports_dir / "chart.svg"

            chart_artifact = next(artifact for artifact in manifest["artifacts"] if artifact["artifact_type"] == "chart")
            self.assertEqual("chart_svg", chart_artifact["artifact_id"])
            self.assertEqual("svg", chart_artifact["format"])
            self.assertEqual("image/svg+xml", chart_artifact["mime_type"])
            self.assertEqual("chart.svg", chart_artifact["file_name"])
            self.assertTrue(chart_path.read_text(encoding="utf-8").startswith("<svg"))
            self.assertFalse((exports_dir / "chart.png").exists())


if __name__ == "__main__":
    unittest.main()
