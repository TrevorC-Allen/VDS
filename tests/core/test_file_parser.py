"""Tests for uploaded file parsing fallbacks."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
import pandas as pd

try:
    import duckdb
except ImportError:  # pragma: no cover - environment dependent.
    duckdb = None

from data_agent_core.core.file_parser import parse_dataset_file


class FileParserTest(unittest.TestCase):
    def test_excel_parser_detects_table_candidates_and_source_roles(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook_path = Path(temp_dir) / "复杂业务表.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "销售明细"
            sheet["A1"] = "销售数据导出"
            sheet["A3"] = "门店信息"
            sheet.merge_cells("A3:B3")
            sheet["C3"] = "财务指标"
            sheet.merge_cells("C3:D3")
            sheet.append([])
            sheet["A4"] = "city"
            sheet["B4"] = "product"
            sheet["C4"] = "sales"
            sheet["D4"] = "profit"
            sheet["A5"] = "上海"
            sheet["B5"] = "苹果"
            sheet["C5"] = 100
            sheet["D5"] = 40
            sheet["A6"] = "北京"
            sheet["B6"] = "香蕉"
            sheet["C6"] = 200
            sheet["D6"] = 60
            rule = workbook.create_sheet("规则说明")
            rule["A1"] = "业务口径"
            rule["A2"] = "利润率=sum利润/sum销售"
            dictionary = workbook.create_sheet("字段字典")
            dictionary.append(["字段名", "含义"])
            dictionary.append(["sales", "销售额"])
            workbook.save(workbook_path)

            parsed = parse_dataset_file(workbook_path, dataset_id="ds_excel")

        self.assertEqual(["销售明细__table_2"], list(parsed.tables))
        table = parsed.tables["销售明细__table_2"]
        self.assertEqual((2, 4), table.shape)
        self.assertIn("财务指标_sales", table.columns)
        self.assertTrue(pd.api.types.is_numeric_dtype(table["财务指标_sales"]))
        self.assertTrue(pd.api.types.is_numeric_dtype(table["财务指标_profit"]))
        roles = {profile.table_name: profile.table_role for profile in parsed.profile.tables}
        self.assertEqual("data_table", roles["销售明细__table_2"])
        self.assertEqual("rule_or_notes", roles["规则说明"])
        self.assertEqual("field_dictionary", roles["字段字典"])
        data_profile = next(profile for profile in parsed.profile.tables if profile.table_name == "销售明细__table_2")
        self.assertEqual("C3:D6".split(":")[-1], data_profile.range_ref.split(":")[-1])
        self.assertIn(4, data_profile.header_rows)
        self.assertIn("treated as rules/notes/metadata", " ".join(parsed.profile.warnings))

    def test_csv_parser_records_encoding_delimiter_and_bad_lines(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            csv_path = Path(temp_dir) / "gbk_sales.csv"
            csv_path.write_bytes("city;sales\n上海;100\nbad;line;extra\n北京;200\n".encode("gbk"))

            parsed = parse_dataset_file(csv_path, dataset_id="ds_csv")

        table = parsed.tables["gbk_sales"]
        self.assertEqual(["city", "sales"], list(table.columns))
        self.assertEqual(2, len(table))
        diagnostics = parsed.profile.parse_diagnostics["sources"][0]
        self.assertEqual(";", diagnostics["delimiter"])
        self.assertEqual("gb18030", diagnostics["encoding"])
        self.assertEqual(1, diagnostics["bad_line_count"])
        self.assertIn("unexpected field counts", " ".join(parsed.profile.warnings))

    def test_image_dataset_upload_boundary_requires_ocr_conversion(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            image_path = Path(temp_dir) / "table.png"
            image_path.write_bytes(b"not really an image")

            with self.assertRaisesRegex(ValueError, "needs OCR"):
                parse_dataset_file(image_path, dataset_id="ds_image")

    @unittest.skipIf(duckdb is None, "duckdb is required to create parquet fixtures")
    def test_parquet_upload_uses_duckdb_when_pandas_engine_is_missing(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            parquet_path = Path(temp_dir) / "taxi_sample.parquet"
            with duckdb.connect(database=":memory:") as con:
                con.execute(
                    f"""
                    COPY (
                        SELECT 1 AS VendorID, 10.5 AS total_amount
                        UNION ALL
                        SELECT 2 AS VendorID, 20.0 AS total_amount
                    ) TO '{parquet_path.as_posix()}' (FORMAT PARQUET)
                    """
                )

            with patch("pandas.read_parquet", side_effect=ImportError("Missing optional dependency 'pyarrow'.")):
                parsed = parse_dataset_file(parquet_path, dataset_id="ds_parquet")

        self.assertEqual(["taxi_sample"], list(parsed.tables))
        table = parsed.tables["taxi_sample"]
        self.assertEqual((2, 2), table.shape)
        self.assertEqual(30.5, float(table["total_amount"].sum()))
        self.assertEqual("ready", parsed.profile.status)


if __name__ == "__main__":
    unittest.main()
